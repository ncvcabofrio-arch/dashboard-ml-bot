"""
INVENTARIO — a foto completa para a contagem fisica. SO LE, nao grava nada.

Junta num arquivo so' as duas metades que faltavam conversar:

  ESTOQUE   getInventoryProductsList + getInventoryProductsStock
            o que o sistema diz que tem, e o que esta reservado

  PEDIDOS   getOrders
            o que ja foi vendido e ainda nao saiu da sua prateleira

O problema que isto resolve
---------------------------
Produto vendido com entrega para depois continua na prateleira, mas ja tem
dono. Se a Base ja baixou do estoque na hora da venda, o conferente vai
achar MAIS peca do que o sistema diz — e se "corrigir" pelo que contou,
recoloca no estoque unidade que ja e' de outro cliente.

A aba CONTAGEM resolve isso pondo tudo na mesma linha:
    estoque_sistema  +  a_entregar  =  esperado_na_prateleira

ATENCAO A UMA HIPOTESE: essa soma vale se a Base baixa o estoque no momento
da VENDA. Se ela so' baixa no DESPACHO, o esperado e' o proprio
estoque_sistema e o a_entregar ja esta dentro dele. Confira em UM SKU com
venda a entregar antes de sair contando o galpao inteiro — a planilha traz
as duas colunas justamente para voce poder decidir olhando.

Variaveis
   BASELINKER_TOKEN            obrigatorio
   BL_INVENTORY_ID             catalogo (padrao: o marcado como default)
   BL_WAREHOUSES               depositos a somar; vazio = todos
   BL_DIAS                     dias de pedidos para tras (padrao 60)
   BL_STATUS_A_ENTREGAR        ids de status que contam como "ainda comigo"
   BL_INCLUIR_NAO_CONFIRMADOS  1 = traz tambem pedidos nao confirmados
   BL_SAIDA                    nome do arquivo (padrao inventario.xlsx)
"""

import json
import os
import sys
import time
from datetime import datetime, timedelta, timezone

import requests

API = "https://api.baselinker.com/connector.php"
TOKEN = (os.environ.get("BASELINKER_TOKEN") or "").strip()
INVENTORY = (os.environ.get("BL_INVENTORY_ID") or "").strip()
SAIDA = os.environ.get("BL_SAIDA", "inventario.xlsx")
DIAS = int(os.environ.get("BL_DIAS", "60") or 60)
NAO_CONFIRMADOS = os.environ.get("BL_INCLUIR_NAO_CONFIRMADOS", "0") == "1"
WAREHOUSES = [w.strip() for w in os.environ.get("BL_WAREHOUSES", "").split(",") if w.strip()]
A_ENTREGAR = [s.strip() for s in
              os.environ.get("BL_STATUS_A_ENTREGAR", "").split(",") if s.strip()]

if not TOKEN:
    print("[ERRO] Falta o BASELINKER_TOKEN.")
    sys.exit(1)


# --------------------------------------------------------------------- API

def chamar(metodo, parametros=None, tentativas=4):
    """A BaseLinker responde HTTP 200 mesmo em erro — quem manda e' o campo
    'status' do corpo. Por isso a checagem e' aqui, nao no status HTTP."""
    corpo = {"method": metodo,
             "parameters": json.dumps(parametros or {}, ensure_ascii=False)}
    for t in range(tentativas):
        try:
            r = requests.post(API, headers={"X-BLToken": TOKEN}, data=corpo, timeout=90)
        except requests.RequestException as e:
            if t == tentativas - 1:
                raise
            print(f"   rede falhou ({e}); tentando de novo...")
            time.sleep(2 * (t + 1))
            continue
        if r.status_code == 429 or r.status_code >= 500:
            time.sleep(3 * (t + 1))
            continue
        d = r.json()
        if d.get("status") == "ERROR":
            code = d.get("error_code", "")
            if code in ("ERROR_RATE_LIMIT", "ERROR_INTERNAL") and t < tentativas - 1:
                time.sleep(3 * (t + 1))
                continue
            raise RuntimeError(f"{metodo}: {code} — {d.get('error_message')}")
        return d
    raise RuntimeError(f"{metodo}: sem resposta")


def num(v):
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def limpo(n):
    n = float(n)
    return int(n) if n.is_integer() else round(n, 2)


def somar(mapa):
    total = 0.0
    for wid, qtd in (mapa or {}).items():
        if WAREHOUSES and wid not in WAREHOUSES:
            continue
        total += num(qtd)
    return total


# ----------------------------------------------------------------- estoque

def inventario_id():
    if INVENTORY:
        return int(INVENTORY)
    invs = chamar("getInventories").get("inventories", []) or []
    if not invs:
        raise RuntimeError("nenhum catalogo encontrado nesta conta")
    padrao = next((i for i in invs if i.get("is_default")), invs[0])
    return int(padrao["inventory_id"])


def paginar(metodo, inv_id, rotulo):
    """Os dois metodos de catalogo paginam por 'page', 1000 por vez."""
    out, page = {}, 1
    while True:
        d = chamar(metodo, {"inventory_id": inv_id, "page": page})
        prods = d.get("products") or {}
        if not prods:
            break
        out.update(prods)
        print(f"  {rotulo}: pagina {page} -> {len(prods)} itens (acumulado {len(out)})")
        if len(prods) < 1000:
            break
        page += 1
        time.sleep(0.65)
    return out


def reservas_por_id(estoques):
    """Achata getInventoryProductsStock em {id: {deposito: qtd}}, incluindo
    variantes — que vem aninhadas no pai, em 'variants' ou em
    'variant_reservations' conforme a versao da conta."""
    mapa = {}
    for pid, p in (estoques or {}).items():
        mapa[str(pid)] = p.get("reservations") or {}
        aninhadas = p.get("variants") or p.get("variant_reservations") or {}
        if isinstance(aninhadas, dict):
            for vid, v in aninhadas.items():
                if isinstance(v, dict) and "reservations" in v:
                    mapa[str(vid)] = v.get("reservations") or {}
                elif isinstance(v, dict):
                    mapa[str(vid)] = v
    return mapa


# ----------------------------------------------------------------- pedidos

def status_da_conta():
    d = chamar("getOrderStatusList")
    return {str(s.get("id")): (s.get("name") or s.get("name_for_customer") or "?")
            for s in (d.get("statuses") or [])}


def baixar_pedidos(desde_ts):
    """getOrders devolve no maximo 100 por vez; a paginacao e' pelo
    date_confirmed do ultimo pedido + 1 segundo.

    Duas armadilhas tratadas: se os 100 tiverem o MESMO date_confirmed o
    cursor nao anda e viraria laco infinito (por isso o +1 forcado); e
    pedido repetido entre paginas contaria unidade duas vezes (por isso o
    controle de ja vistos)."""
    pedidos, vistos, cursor, paginas = [], set(), int(desde_ts), 0
    while True:
        params = {"date_confirmed_from": int(cursor)}
        if NAO_CONFIRMADOS:
            params["get_unconfirmed_orders"] = True
        lote = chamar("getOrders", params).get("orders", []) or []
        novos = [o for o in lote if str(o.get("order_id")) not in vistos]
        for o in novos:
            vistos.add(str(o.get("order_id")))
        pedidos.extend(novos)
        paginas += 1
        print(f"  pedidos: pagina {paginas} -> {len(lote)} recebidos, "
              f"{len(novos)} novos (acumulado {len(pedidos)})")
        if len(lote) < 100:
            break
        cursor = max(max(int(num(o.get("date_confirmed"))) for o in lote) + 1,
                     cursor + 1)
        if paginas >= 300:
            print("  ⚠️ parei em 300 paginas por seguranca.")
            break
        time.sleep(0.65)
    return pedidos


# -------------------------------------------------------------------- main

def main():
    inv_id = inventario_id()
    print(f"Catalogo {inv_id} | depositos: "
          f"{', '.join(WAREHOUSES) if WAREHOUSES else 'TODOS'}")

    print("\n[1/2] ESTOQUE")
    lista = paginar("getInventoryProductsList", inv_id, "lista")
    reservas = {}
    try:
        reservas = reservas_por_id(paginar("getInventoryProductsStock", inv_id, "estoque"))
    except Exception as e:
        print(f"  ⚠️ nao consegui ler as reservas ({e}); coluna vai zerada.")

    estoque, reservado, nome_cat, extra = {}, {}, {}, {}
    for pid, p in lista.items():
        sku = str(p.get("sku") or "").strip()
        if not sku:
            continue
        estoque[sku] = estoque.get(sku, 0.0) + somar(p.get("stock"))
        reservado[sku] = reservado.get(sku, 0.0) + somar(reservas.get(str(pid)))
        nome_cat.setdefault(sku, p.get("name") or "")
        extra.setdefault(sku, (pid, p.get("ean") or "",
                               1 if p.get("parent_id") else 0))

    print("\n[2/2] PEDIDOS")
    desde = datetime.now(timezone.utc) - timedelta(days=DIAS)
    print(f"  confirmados desde {desde:%d/%m/%Y} ({DIAS} dias)"
          + ("  [incluindo nao confirmados]" if NAO_CONFIRMADOS else ""))

    nomes = {}
    try:
        nomes = status_da_conta()
        print(f"  status cadastrados na conta: {len(nomes)}")
    except Exception as e:
        print(f"  aviso: nao li a lista de status ({e}); mostro so' os ids.")

    pedidos = baixar_pedidos(desde.timestamp())

    por_status, matriz, etiquetados, nome_ped, detalhe = {}, {}, {}, {}, []
    if pedidos:
        print("\n  Campos que os SEUS pedidos trazem (para conferencia):")
        print("  " + ", ".join(sorted(pedidos[0].keys())))

    for o in pedidos:
        sid = str(o.get("order_status_id"))
        reg = por_status.setdefault(sid, {"pedidos": 0, "unidades": 0.0,
                                          "com_etiqueta": 0})
        reg["pedidos"] += 1

        # ETIQUETA: a Base preenche o numero da encomenda quando ela e' gerada.
        # Vazio = nao imprimiu, entao a peca continua no galpao.
        nr = str(o.get("delivery_package_nr") or "").strip()
        tem_etiqueta = bool(nr) or bool(str(o.get("delivery_package_module") or "").strip())
        if tem_etiqueta:
            reg["com_etiqueta"] += 1

        quando = o.get("date_confirmed") or o.get("date_add") or 0
        try:
            data = datetime.fromtimestamp(int(num(quando)), timezone.utc).strftime("%d/%m/%Y")
        except (ValueError, OSError):
            data = ""

        for p in o.get("products", []) or []:
            sku = str(p.get("sku") or "").strip()
            qtd = num(p.get("quantity"))
            reg["unidades"] += qtd
            if not sku:
                continue
            matriz.setdefault(sku, {})
            matriz[sku][sid] = matriz[sku].get(sid, 0.0) + qtd
            if tem_etiqueta:
                etiquetados.setdefault(sku, {})
                etiquetados[sku][sid] = etiquetados[sku].get(sid, 0.0) + qtd
            nome_ped.setdefault(sku, p.get("name") or "")
            detalhe.append([o.get("order_id"), data, nomes.get(sid, f"status {sid}"),
                            sid, "sim" if tem_etiqueta else "nao", nr,
                            o.get("delivery_method") or "", o.get("pick_state", ""),
                            o.get("pack_state", ""), sku, limpo(qtd),
                            p.get("name") or ""])

    ordenados = sorted(por_status.items(), key=lambda x: -x[1]["pedidos"])

    # ---------------------------------------------------------- console
    print("\n" + "=" * 70)
    print("PEDIDOS POR STATUS")
    print("=" * 70)
    print(f"  {'id':>6}  {'status':<32} {'pedidos':>8} {'unidades':>9} {'c/etiqueta':>11}")
    for sid, v in ordenados:
        print(f"  {sid:>6}  {nomes.get(sid, '(sem nome)'):<32} "
              f"{v['pedidos']:>8} {limpo(v['unidades']):>9} {v['com_etiqueta']:>11}")
    if ordenados:
        print("\n  'c/etiqueta' = pedidos com numero de encomenda ja gerado.")
        print("  Status com muitos pedidos e ZERO etiqueta e' candidato a")
        print("  'ainda na prateleira'. Status todo etiquetado ja saiu.")

    alvo = set(A_ENTREGAR)
    if alvo:
        print(f"\n  Contando como A ENTREGAR: "
              + ", ".join(f"{s} ({nomes.get(s, '?')})" for s in A_ENTREGAR))
    elif ordenados:
        print("\n  BL_STATUS_A_ENTREGAR nao configurado — esta rodada e' de")
        print("  reconhecimento. Escolha na tabela acima os status que")
        print("  significam 'vendido mas ainda comigo' e rode de novo com")
        print("  esses ids: ai a aba CONTAGEM sai preenchida.")

    def a_entregar_de(sku):
        return sum(q for s, q in matriz.get(sku, {}).items() if s in alvo)

    def etiquetadas_de(sku):
        return sum(q for s, q in etiquetados.get(sku, {}).items() if s in alvo)

    # ---------------------------------------------------------- planilha
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("\n(openpyxl nao instalado; sem planilha)")
        return

    wb = Workbook()

    def cabecalho(ws, cols, larguras, cor="DDDDDD"):
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor=cor)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"
        for i, w in enumerate(larguras, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # ---- CONTAGEM: a aba que vai para a mao do conferente
    ws = wb.active
    ws.title = "Contagem"
    cabecalho(ws, ["sku", "nome", "estoque_sistema", "reservado", "a_entregar",
                   "dessas_etiquetadas", "esperado_na_prateleira", "contado",
                   "diferenca"],
              [22, 55, 16, 11, 12, 18, 21, 10, 11], "FFF2CC")
    todos = sorted(set(estoque) | set(matriz))
    destaque = PatternFill("solid", fgColor="FFE08A")
    for sku in todos:
        ae = a_entregar_de(sku)
        est = estoque.get(sku, 0.0)
        linha = ws.max_row + 1
        ws.append([sku, nome_cat.get(sku) or nome_ped.get(sku, ""),
                   limpo(est), limpo(reservado.get(sku, 0.0)),
                   limpo(ae), limpo(etiquetadas_de(sku)), limpo(est + ae),
                   None, f"=IF(H{linha}=\"\",\"\",H{linha}-G{linha})"])
        if ae > 0:
            for c in ws[linha]:
                c.fill = destaque
    ws.auto_filter.ref = f"A1:I{ws.max_row}"

    # ---- estoque cru
    ws2 = wb.create_sheet("Estoque x Reserva")
    cabecalho(ws2, ["sku", "nome", "estoque", "reservado", "disponivel",
                    "product_id", "ean", "e_variante"],
              [22, 65, 10, 11, 12, 13, 16, 11])
    for sku in sorted(estoque):
        pid, ean, var = extra.get(sku, ("", "", 0))
        est, res = estoque[sku], reservado.get(sku, 0.0)
        ws2.append([sku, nome_cat.get(sku, ""), limpo(est), limpo(res),
                    limpo(est - res), pid, ean, var])

    # ---- pedidos
    if ordenados:
        ws3 = wb.create_sheet("Resumo por status")
        cabecalho(ws3, ["status_id", "status", "pedidos", "unidades",
                        "pedidos_com_etiqueta"], [12, 40, 10, 11, 20])
        for sid, v in ordenados:
            ws3.append([sid, nomes.get(sid, ""), v["pedidos"],
                        limpo(v["unidades"]), v["com_etiqueta"]])

        ids = [sid for sid, _ in ordenados]
        ws4 = wb.create_sheet("SKU x status")
        cabecalho(ws4, ["sku", "nome"] + [nomes.get(s, f"status {s}") for s in ids]
                  + ["total"], [22, 55] + [16] * len(ids) + [10])
        for sku in sorted(matriz, key=lambda s: -sum(matriz[s].values())):
            m = matriz[sku]
            ws4.append([sku, nome_ped.get(sku, "")]
                       + [limpo(m.get(s, 0)) for s in ids]
                       + [limpo(sum(m.values()))])

        ws5 = wb.create_sheet("Pedidos")
        cabecalho(ws5, ["order_id", "data", "status", "status_id", "etiqueta",
                        "nr_encomenda", "metodo_envio", "pick_state",
                        "pack_state", "sku", "quantidade", "produto"],
                  [16, 12, 26, 10, 10, 22, 22, 11, 11, 20, 11, 50])
        for l in detalhe:
            ws5.append(l)

    wb.save(SAIDA)

    # ---------------------------------------------------------- resumo
    com_est = [s for s in estoque if estoque[s] > 0]
    com_res = [s for s in reservado if reservado[s] > 0]
    com_ae = [s for s in todos if a_entregar_de(s) > 0]
    print("\n" + "=" * 70)
    print("RESUMO")
    print("=" * 70)
    print(f"  SKUs no catalogo .............. {len(estoque)}")
    print(f"  com estoque ................... {len(com_est)} "
          f"({limpo(sum(estoque[s] for s in com_est))} unidades)")
    print(f"  com reserva na Base ........... {len(com_res)}")
    print(f"  com venda a entregar .......... {len(com_ae)} "
          f"({limpo(sum(a_entregar_de(s) for s in com_ae))} unidades)")
    print(f"\n✅ {SAIDA}")
    print("   Aba CONTAGEM: preencha a coluna 'contado' e a diferenca sai sozinha.")


if __name__ == "__main__":
    main()
